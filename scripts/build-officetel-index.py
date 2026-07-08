#!/usr/bin/env python3
"""국세청 상업용건물/오피스텔 기준시가 XLSX → 시도별 파셀 인덱스 JSON 빌드.

사용법:
  python3 scripts/build-officetel-index.py <기준시가.xlsx> [출력디렉토리=data/officetel-prices]

입력: 공공데이터포털 "국세청_상업용건물 오피스텔 기준시가" XLSX (시트 5개, 각 50만 행).
출력: data/officetel-prices/{법정동코드 앞2자리}.json.gz (gzip)
  { "parcels": { "법정동코드10:번지4:호4": "유형|동|층|호|단가|전용|공유;..." } }
  - 유형: O=오피스텔, S=상가
  - 층: 지하는 B 접두(B1), 지상은 숫자만
기준시가 총액 = 단가(원/㎡) × (전용면적 + 공유면적)

의존성: pip install openpyxl
"""
import gzip
import json
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    source = Path(sys.argv[1])
    out_dir = Path(sys.argv[2] if len(sys.argv) > 2 else "data/officetel-prices")
    out_dir.mkdir(parents=True, exist_ok=True)

    shards = defaultdict(lambda: defaultdict(list))
    counts = defaultdict(int)
    skipped = 0
    started = time.time()

    wb = openpyxl.load_workbook(source, read_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[3] is None:
                continue
            kind_raw = str(row[1] or "")
            special = str(row[4] or "")
            if special and special != "일반지번":
                skipped += 1
                continue

            ld_code = str(row[3]).zfill(10)
            bunji = str(row[5] or "").zfill(4)
            ho_sub = str(row[6] or "").zfill(4)
            dong = str(row[8] or "").replace("(단일)", "").strip()
            if dong in ("1", "0"):  # 단일동 표기는 비워서 동 무관 매칭
                dong = ""
            floor_type = str(row[9] or "")
            floor = str(row[10] or "").strip()
            unit_ho = str(row[11] or "").strip()
            price = row[12]
            excl = row[13]
            share = row[14]
            if not unit_ho or price is None:
                skipped += 1
                continue

            kind = "O" if "오피스텔" in kind_raw else "S"
            counts[kind] += 1
            floor_label = ("B" + floor) if "지하" in floor_type else floor
            record = f"{kind}|{dong}|{floor_label}|{unit_ho}|{price}|{excl or 0}|{share or 0}"
            key = f"{ld_code}:{bunji}:{ho_sub}"
            shards[ld_code[:2]][key].append(record)

    total = sum(counts.values())
    for sido, parcels in sorted(shards.items()):
        payload = {
            "source": "국세청 상업용건물/오피스텔 기준시가",
            "generatedAt": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "parcels": {key: ";".join(records) for key, records in parcels.items()},
        }
        path = out_dir / f"{sido}.json.gz"
        with gzip.open(path, "wt", encoding="utf-8", compresslevel=9) as handle:
            json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
        print(f"  {path} — 파셀 {len(parcels):,}")

    print(f"완료: 세대 {total:,} (오피스텔 {counts['O']:,} / 상가 {counts['S']:,}), 제외 {skipped:,}, {time.time()-started:.0f}초")


if __name__ == "__main__":
    main()
